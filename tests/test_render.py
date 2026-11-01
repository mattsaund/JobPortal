import asyncio
from datetime import date

import pytest

from jobportal.core import osutil
from jobportal.documents import render, resumecheck


def test_resume_html(sample_resume):
    html = render.resume_html(sample_resume)
    assert "<h1>Jordan Rivera</h1>" in html
    assert '<div class="line split">' in html  # "left || dates" puts the dates flush right
    assert "<li>" in html and "<script" not in html


def test_links_are_sanitized():
    html = render.inline_html("[click](javascript:alert(1)) and [ok](https://example.com)")
    assert 'href="#"' in html and 'href="https://example.com"' in html


def test_bare_links_stay_inside_their_tags():
    # A bold address is one link inside the bold, not a broken tag.
    assert render.inline_html("**https://example.com**") == '<b><a href="https://example.com">https://example.com</a></b>'
    # A quote ends the address, so text after it can't add attributes to the link.
    assert render.inline_html('see https://example.com"onmouseover="alert(1)').startswith('see <a href="https://example.com">')
    # Punctuation that ends a sentence stays out of the link.
    assert 'href="https://example.com/portfolio">' in render.inline_html("My work: https://example.com/portfolio.")


def test_docx_links_and_numbered_lines(tmp_path):
    from docx import Document
    out = tmp_path / "prep.docx"
    md = "## Steps\n1. First\n2. Second\n- Read [the guide](https://example.com/guide), not [this](javascript:alert(1))\n- **[Bold link](https://example.com/b)**"
    render.markdown_docx(md, out, "Interview prep")
    doc = Document(str(out))
    assert "1. First" in [p.text for p in doc.paragraphs]  # numbers print as written; Word never renumbers them
    links = [link for p in doc.paragraphs for link in p.hyperlinks]
    assert [link.url for link in links] == ["https://example.com/guide", "https://example.com/b"]  # the javascript: link is plain text
    assert links[1].runs[0].bold


def test_control_characters_never_break_a_word_file(tmp_path):
    from docx import Document
    # \x0b is the line break text pasted from Word carries; python-docx used to refuse the whole file over it.
    render.resume_docx("# Jordan Rivera\njordan.rivera@example.com\n\n## Experience\n- Line one\x0bline two\x01", tmp_path / "r.docx")
    assert any("Line one line two" in p.text for p in Document(str(tmp_path / "r.docx")).paragraphs)
    render.cover_docx("Dear team,\x0b\n\nThank you.", {"name": "Jordan\x0bRivera", "links": {"site": "example.com\x02"}},
                      "Example\x0c Co", "", tmp_path / "c.docx")
    render.markdown_docx("## Prep\x07\n- Ask about the team\x0b", tmp_path / "p.docx", "Interview prep\x0b", "Example\x1f Co")
    assert all((tmp_path / name).stat().st_size > 1000 for name in ("r.docx", "c.docx", "p.docx"))
    assert b"Subject: Following up on the Data Intern role" in render.eml("hr@example.com", "Following up\non the Data Intern role", "Hi")


def test_the_fit_loop_never_leaves_an_overflowing_try_in_place(tmp_path, monkeypatch):
    out, tries = tmp_path / "resume.pdf", []

    async def fake_print(html_text, path):
        assert not out.exists()  # a download during the loop must not find a try that runs to two pages
        path.write_bytes(b"%PDF-1.4 try")
        tries.append(path)
        return 2 if len(tries) < 3 else 1  # the third scale fits on one page

    monkeypatch.setattr(render, "html_to_pdf", fake_print)
    info = asyncio.run(render.resume_pdf("# Jordan Rivera", out))
    assert info["pages"] == 1 and info["scale"] == render.FIT_SCALES[2] and render.fit_scale(out) == render.FIT_SCALES[2]
    assert sorted(p.name for p in tmp_path.iterdir()) == ["resume.fit", "resume.pdf"]  # no drafts left behind


def test_long_date_is_portable():
    assert render.long_date(date(2026, 3, 5)) == "March 5, 2026"


def test_cover_html_uses_facts():
    facts = {"name": "Jordan Rivera", "email": "jordan.rivera@example.com", "city": "Columbus", "state": "OH", "links": {}}
    html = render.cover_html("Dear Hiring Team,\n\nHello.\n\nSincerely,\nJordan Rivera", facts, "Example Co", "Remote")
    assert "Jordan Rivera" in html and "Example Co" in html


def test_resume_check(sample_resume):
    inv = [{"skill": s, "level": "strong"} for s in ("Python", "JavaScript", "React", "Node.js", "SQL", "PostgreSQL", "Git", "Linux")]
    good = resumecheck.check(sample_resume, inventory=inv, keywords=["Python", "SQL"], pages=1, max_pages=1)
    assert good["score"] >= 85, good["issues"]
    assert not any("phone" in i["text"] for i in good["issues"])  # "(555) 010-0199" counts as a phone number
    bad = resumecheck.check("# Name\nno contact\n\n## My Journey\n- I was responsible for stuff", inventory=inv, pages=2, max_pages=1)
    texts = " ".join(i["text"] for i in bad["issues"])
    assert "2 pages" in texts and "Non-standard heading" in texts and "first person" in texts


def test_short_inventory_skills_only_vouch_for_themselves():
    def unverified(skills_line: str, inventory: list[str]) -> list[str]:
        md = f"# Jordan Rivera\njordan.rivera@example.com\n\n## Skills\n{skills_line}"
        return resumecheck.check(md, inventory=[{"skill": s, "level": "strong"} for s in inventory])["unverified_skills"]

    # Before: R vouched for Docker, Go for Django and Google Cloud, Java for JavaScript, Git for Digital Marketing.
    assert unverified("Docker, Django, Google Cloud, JavaScript, Digital Marketing", ["R", "Go", "Java", "Git", "C"]) == [
        "Docker", "Django", "Google Cloud", "JavaScript", "Digital Marketing"]
    assert unverified("R, Go, React Native, Microsoft Excel, Python (pandas, NumPy)", ["R", "Go", "React", "Excel", "Python"]) == []
    assert unverified("Excel", ["Microsoft Excel"]) == []  # a longer name still covers its last word


def test_roman_numerals_are_not_first_person(sample_resume):
    courses = sample_resume.replace("Operating Systems", "Calculus I, Physics I and II, File I/O")
    assert not any("first person" in i["text"] for i in resumecheck.check(courses)["issues"])
    pronoun = sample_resume.replace("- Resolved", "- I resolved")
    assert any("first person" in i["text"] for i in resumecheck.check(pronoun)["issues"])
    numerals = sample_resume.replace("- Resolved", "- Taught at a Title I school; ran Phase I trials at a Level I trauma center. Resolved")
    addresses = sample_resume.replace("- Resolved", "- Portfolio at about.me/jordan-example, mail hire.me@example.com. Resolved")
    for text in (numerals, addresses):
        assert not any("first person" in i["text"] for i in resumecheck.check(text)["issues"])


@pytest.mark.skipif(not osutil.find_pdf_engine(), reason="needs Chrome, Edge or Chromium")
def test_pdf_export(sample_resume, tmp_path):
    info = asyncio.run(render.resume_pdf(sample_resume, tmp_path / "resume.pdf"))
    assert info["pages"] == 1 and (tmp_path / "resume.pdf").stat().st_size > 1000
    render.resume_docx(sample_resume, tmp_path / "resume.docx", render.fit_scale(tmp_path / "resume.pdf"))
    assert (tmp_path / "resume.docx").stat().st_size > 1000


@pytest.mark.skipif(not osutil.find_pdf_engine(), reason="needs Chrome, Edge or Chromium")
def test_every_pdf_page_keeps_its_margins(tmp_path):
    from pypdf import PdfReader
    md = "\n".join(f"## Topic {i}\n" + "\n".join(f"- Point {j} about topic {i}" for j in range(8)) for i in range(10))
    assert asyncio.run(render.markdown_pdf(md, tmp_path / "prep.pdf", "Interview prep")) >= 2
    page = PdfReader(str(tmp_path / "prep.pdf")).pages[1]
    baselines = []
    page.extract_text(visitor_text=lambda text, cm, tm, font, size: baselines.append(tm[5] * cm[3] + cm[5]) if text.strip() else None)
    assert float(page.mediabox.height) - max(baselines) > 36  # page 2 starts below its top margin, not at the paper's edge
