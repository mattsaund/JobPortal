import io
import zipfile

import pytest

from jobportal.documents import extract


def make_pdf(lines: list[str]) -> bytes:
    """A minimal one-page PDF with real text (Helvetica), built by hand so the test needs no PDF tools."""
    escaped = [ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for ln in lines]
    text = "BT /F1 12 Tf 14 TL 72 740 Td " + " ".join(f"({ln}) '" for ln in escaped) + " ET"
    stream = text.encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


LINES = ["Jordan Rivera", "jordan.rivera@example.com", "Experience", "Student Technician, Campus IT Help Desk",
         "Resolved 40+ hardware and software tickets a week", "Skills: Python, SQL, React"]


def test_pdf_text_and_page_count(tmp_path):
    data = make_pdf(LINES)
    text = extract.extract_text("resume.pdf", data)
    assert "Jordan Rivera" in text and "Campus IT Help Desk" in text
    p = tmp_path / "r.pdf"
    p.write_bytes(data)
    assert extract.pdf_page_count(p) == 1


def test_docx_text_includes_tables_and_header():
    from docx import Document
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Jordan Rivera | jordan.rivera@example.com"
    doc.add_paragraph("Experience")
    doc.add_paragraph("Student Technician at the campus help desk, resolving 40+ tickets every week for students and staff.")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Skills"
    t.cell(0, 1).text = "Python, SQL, React"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract.extract_text("resume.docx", buf.getvalue())
    assert "jordan.rivera@example.com" in text
    assert "Python, SQL, React" in text
    assert text.index("Experience") < text.index("Skills")


def test_docx_text_inside_content_controls_and_header_tables():
    """Word's resume templates wrap text in content controls (w:sdt), and many put the contact details in a header table."""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    doc = Document()
    section = doc.sections[0]
    row = section.header.add_table(rows=1, cols=2, width=section.page_width).rows[0]
    row.cells[0].text, row.cells[1].text = "jordan.rivera@example.com", "(555) 010-0199"
    doc.add_paragraph("Experience")
    block = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>Student Technician, Campus IT Help Desk, resolving '
                      "40+ tickets every week for students and staff.</w:t></w:r></w:p></w:sdtContent></w:sdt>")
    inline = parse_xml(f'<w:p {nsdecls("w")}><w:r><w:t xml:space="preserve">Skills: </w:t></w:r><w:sdt><w:sdtContent><w:r>'
                       "<w:t>Python, SQL</w:t></w:r></w:sdtContent></w:sdt></w:p>")
    doc.paragraphs[-1]._p.addnext(block)
    block.addnext(inline)
    buf = io.BytesIO()
    doc.save(buf)
    text = extract.extract_text("resume.docx", buf.getvalue())
    assert "jordan.rivera@example.com   (555) 010-0199" in text
    assert "Student Technician, Campus IT Help Desk" in text and "Skills: Python, SQL" in text


def test_odt_text():
    content = ('<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
               '<text:h>Jordan Rivera</text:h><text:p>Student Technician<text:tab/>2024 to present</text:p>'
               '<text:p>Resolved hardware and software tickets for students and staff every week &amp; more.</text:p>'
               '</office:text></office:body></office:document-content>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
    text = extract.extract_text("resume.odt", buf.getvalue())
    assert "Jordan Rivera" in text and "Student Technician\t2024" in text and "& more" in text


def test_rtf_text():
    rtf = (r"{\rtf1\ansi\deff0{\fonttbl{\f0 Arial;}}{\colortbl;\red0\green0\blue0;}"
           r"\f0\fs22 Jordan Rivera\par Caf\'e9 shift lead\tab 2023\par "
           r"Resolved tickets for students and staff every week, wrote scripts that saved hours\u8217?s worth of work.\par}")
    text = extract.extract_text("resume.rtf", rtf.encode("latin-1"))
    assert "Jordan Rivera" in text
    assert "Café shift lead\t2023" in text
    assert "’s worth" in text
    assert "Arial" not in text  # the font table is skipped


def test_rtf_from_textedit_and_word_symbols():
    """macOS TextEdit ends lines with a backslash before the line break, and Word writes curly quotes as control words."""
    rtf = ("{\\rtf1\\ansi\\cocoartf2580{\\fonttbl\\f0 Helvetica;}\n\\f0 Jordan Rivera\\\nStudent Technician, Campus IT Help Desk\\\n"
           "Jordan\\rquote s scripts saved the help desk hours of work every week {\\uc1\\u-10179?\\u-8704?}\\par}")
    lines = extract.extract_text("resume.rtf", rtf.encode("latin-1")).split("\n")
    assert lines[:2] == ["Jordan Rivera", "Student Technician, Campus IT Help Desk"]
    assert lines[2].startswith("Jordan’s scripts") and lines[2].endswith("every week \U0001F600")  # two \u halves make one emoji


def test_text_encodings():
    body = "Jordan Rivera\nSoftware engineer who likes résumé parsers and writes Python every day for fun and work."
    assert "résumé" in extract.extract_text("r.txt", body.encode("utf-8"))
    assert "résumé" in extract.extract_text("r.txt", b"\xef\xbb\xbf" + body.encode("utf-8"))
    assert "résumé" in extract.extract_text("r.md", body.encode("cp1252"))


def test_rejects_empty_and_unknown_files():
    with pytest.raises(extract.ExtractError):
        extract.extract_text("r.txt", b"too short")
    with pytest.raises(extract.ExtractError):
        extract.extract_text("photo.png", b"\x89PNG" + b"0" * 200)


def test_a_protected_pdf_gets_a_message_the_user_can_act_on():
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(make_pdf(LINES))))
    writer.encrypt(user_password="secret", algorithm="RC4-128")  # RC4 needs no extra packages to write
    locked = io.BytesIO()
    writer.write(locked)
    with pytest.raises(extract.ExtractError, match="PDF is protected"):  # not "FileNotDecryptedError: ..."
        extract.extract_text("resume.pdf", locked.getvalue())


@pytest.mark.parametrize("exit_code, written", [
    # What LibreOffice writes for a damaged .doc it opens as plain text: the raw bytes, with control bytes shown as "#".
    (0, ("ÐÏ#à¡±#á" + "#" * 16 + ";###þ?\t" + "#" * 30).encode() * 20),
    (1, None),  # the conversion failed and wrote nothing
])
def test_a_doc_libreoffice_cannot_read_is_refused(monkeypatch, exit_code, written):
    import subprocess
    from pathlib import Path

    def fake_soffice(cmd, **kwargs):
        if written is not None:
            (Path(cmd[cmd.index("--outdir") + 1]) / "resume.txt").write_bytes(written)
        return subprocess.CompletedProcess(cmd, exit_code, b"", b"")

    monkeypatch.setattr(extract.osutil, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(extract.subprocess, "run", fake_soffice)
    with pytest.raises(extract.ExtractError, match="LibreOffice couldn't read old.doc"):
        extract.extract_text("old.doc", b"\xd0\xcf\x11\xe0" + b"\x00" * 600)
