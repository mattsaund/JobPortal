"""Resume and cover-letter rendering: markdown to a structured doc, then HTML (preview, and PDF through a headless Chromium
browser such as Chrome or Edge) and DOCX. The classic layout is a conventional one-page resume: Arial, centered name, ruled
section headers, tight bullets."""
import asyncio
import json
import html
import re
import tempfile
from datetime import date
from pathlib import Path

from . import osutil
from .config import CACHE
from .extract import pdf_page_count


def long_date(d: date | None = None) -> str:
    """Dates like "October 5, 2026" on every OS (strftime's %-d doesn't exist on Windows)."""
    d = d or date.today()
    return f"{d:%B} {d.day}, {d.year}"

# ---------------------------------------------------------------- parsing

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")
URL_RE = re.compile(r"(?<![\w/\"'=])(https?://[^\s,)<]+)")


def safe_href(u: str) -> str:
    """Only http(s)/mailto links survive: AI output and job posts come from the open web."""
    u = (u or "").strip()
    return u if re.match(r"^(https?://|mailto:)", u, re.I) else "#"


def inline_html(s: str) -> str:
    links = []

    def keep(m):
        links.append((m.group(1), safe_href(m.group(2))))
        return f"\x00{len(links) - 1}\x00"

    s = LINK_RE.sub(keep, s)
    s = html.escape(s, quote=False)
    s = URL_RE.sub(lambda m: f'<a href="{m.group(1)}">{m.group(1)}</a>', s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])", r"<i>\1</i>", s)
    return re.sub(r"\x00(\d+)\x00", lambda m: f'<a href="{html.escape(links[int(m.group(1))][1])}">{html.escape(links[int(m.group(1))][0])}</a>', s)


def parse_resume(md: str) -> dict:
    doc = {"name": "", "contact": [], "sections": []}
    sec = None
    entry: list | None = None
    for raw in (md or "").replace("\r", "").split("\n"):
        line = raw.rstrip()
        st = line.strip()
        if st.startswith("# ") and not doc["name"]:
            doc["name"] = st[2:].strip()
            continue
        if st.startswith("## "):
            sec = {"title": st[3:].strip(), "entries": []}
            doc["sections"].append(sec)
            entry = None
            continue
        if not st:
            entry = None
            continue
        if sec is None:
            doc["contact"].append(st)
            continue
        if entry is None:
            entry = []
            sec["entries"].append(entry)
        if re.match(r"^[-*•]\s+", st):
            entry.append(("bullet", re.sub(r"^[-*•]\s+", "", st)))
        else:
            entry.append(("line", st.lstrip("#").strip()))
    return doc


# ---------------------------------------------------------------- HTML

RESUME_BASE_CSS = """
@page { size: Letter; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; }
body { font-family: Arial, "Liberation Sans", Helvetica, sans-serif; color: #000; -webkit-print-color-adjust: exact; }
.page { width: 8.5in; min-height: 11in; padding: var(--mt, 0.45in) var(--mx, 0.5in) 0.4in; font-size: calc(10pt * var(--s, 1)); line-height: 1.17; }
h1 { font-size: calc(16.5pt * var(--s, 1)); text-align: center; margin: 0 0 3pt; font-weight: 700; }
.contact { text-align: center; font-size: calc(11pt * var(--s, 1)); margin-bottom: 2pt; }
a { color: #1155cc; text-decoration: underline; }
h2 { font-size: calc(11.5pt * var(--s, 1)); font-weight: 700; margin: calc(7pt * var(--s, 1)) 0 calc(2pt * var(--s, 1));
     border-bottom: 1px solid #000; padding-bottom: 1pt; }
.entry { margin: 0 0 calc(5pt * var(--s, 1)); }
.entry:last-child { margin-bottom: 0; }
.line { margin: 0; }
.line.split { display: flex; justify-content: space-between; gap: 12pt; }
.line.split .r { white-space: nowrap; }
ul { margin: 0; padding-left: 0.26in; }
li { margin: 0; padding-left: 0.04in; }
"""

# Every layout is single-column, text-only and ATS-safe; they differ only in typography and header alignment.
RESUME_THEMES = {
    "classic": "",
    "modern": """
h1 { text-align: left; font-size: calc(19pt * var(--s, 1)); margin-bottom: 2pt; letter-spacing: .2pt; }
.contact { text-align: left; font-size: calc(10pt * var(--s, 1)); color: #222; }
h2 { font-size: calc(10.5pt * var(--s, 1)); text-transform: uppercase; letter-spacing: .9pt; margin-top: calc(8pt * var(--s, 1)); }
a { color: #000; text-decoration: none; }
""",
    "compact": """
.page { font-size: calc(9.6pt * var(--s, 1)); line-height: 1.12; }
h1 { font-size: calc(15pt * var(--s, 1)); margin-bottom: 1pt; }
.contact { font-size: calc(10pt * var(--s, 1)); }
h2 { font-size: calc(10.5pt * var(--s, 1)); margin: calc(5pt * var(--s, 1)) 0 calc(1.5pt * var(--s, 1)); }
.entry { margin-bottom: calc(3.5pt * var(--s, 1)); }
""",
    "academic": """
body { font-family: "Times New Roman", "Liberation Serif", Times, serif; }
.page { font-size: calc(10.8pt * var(--s, 1)); line-height: 1.15; }
h1 { font-size: calc(18pt * var(--s, 1)); font-weight: 700; }
h2 { font-size: calc(12pt * var(--s, 1)); font-variant: small-caps; letter-spacing: .4pt; }
""",
}
THEMES = list(RESUME_THEMES)


COVER_CSS = """
@page { size: Letter; margin: 0; }
html, body { margin: 0; padding: 0; background: #fff; }
body { font-family: Arial, "Liberation Sans", Helvetica, sans-serif; color: #000; }
.page { width: 8.5in; min-height: 11in; padding: 0.9in 1in; font-size: calc(11pt * var(--s, 1)); line-height: 1.42; }
.hd h1 { font-size: 18pt; margin: 0 0 2pt; }
.hd .contact { font-size: 10pt; color: #222; }
.hd { border-bottom: 1px solid #000; padding-bottom: 8pt; margin-bottom: 18pt; }
.meta { margin-bottom: 14pt; }
p { margin: 0 0 11pt; }
a { color: #1155cc; }
"""


def _line_html(text: str) -> str:
    """A resume line; "left || right" puts the right part (usually dates) flush right on the same line."""
    if " || " in text:
        left, right = text.split(" || ", 1)
        return f'<div class="line split"><span>{inline_html(left)}</span><span class="r">{inline_html(right)}</span></div>'
    return f'<div class="line">{inline_html(text)}</div>'


def resume_html(md: str, scale: float = 1.0, theme: str = "classic") -> str:
    d = parse_resume(md)
    parts = [f"<h1>{inline_html(d['name'])}</h1>"]
    parts += [f'<div class="contact">{inline_html(c)}</div>' for c in d["contact"]]
    for s in d["sections"]:
        parts.append(f"<h2>{inline_html(s['title'])}</h2>")
        for e in s["entries"]:
            buf, in_ul = ['<div class="entry">'], False
            for kind, text in e:
                if kind == "bullet" and not in_ul:
                    buf.append("<ul>")
                    in_ul = True
                if kind != "bullet" and in_ul:
                    buf.append("</ul>")
                    in_ul = False
                buf.append(f"<li>{inline_html(text.replace(' || ', ' | '))}</li>" if kind == "bullet" else _line_html(text))
            if in_ul:
                buf.append("</ul>")
            buf.append("</div>")
            parts.append("".join(buf))
    mt = 0.45 if scale >= 0.95 else 0.35
    mx = 0.47 if scale >= 0.95 else 0.42
    css = RESUME_BASE_CSS + RESUME_THEMES.get(theme, "")
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(d["name"])} Resume</title>'
            f"<style>{css}</style></head><body><div class=\"page\" style=\"--s:{scale};--mt:{mt}in;--mx:{mx}in\">"
            + "\n".join(parts) + "</div></body></html>")


def cover_html(md: str, facts: dict, company: str = "", location: str = "", scale: float = 1.0) -> str:
    links = facts.get("links", {})
    contact = " | ".join(x for x in [
        facts.get("email"), facts.get("phone"), f"{facts.get('city', '')}, {facts.get('state', '')}".strip(", "),
        (links.get("site") or "").replace("https://", ""), (links.get("linkedin") or "").replace("https://www.", "")] if x)
    paras = [p.strip() for p in re.split(r"\n\s*\n", (md or "").strip()) if p.strip()]
    body = "".join(f"<p>{'<br>'.join(inline_html(l) for l in p.split(chr(10)))}</p>" for p in paras)
    meta = long_date()
    if company:
        meta += f"<br><br>Hiring Team<br>{html.escape(company)}" + (f"<br>{html.escape(location)}" if location else "")
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>Cover Letter</title><style>{COVER_CSS}</style></head>'
            f'<body><div class="page" style="--s:{scale}"><div class="hd"><h1>{html.escape(facts.get("name", ""))}</h1>'
            f'<div class="contact">{inline_html(contact)}</div></div><div class="meta">{meta}</div>{body}</div></body></html>')


# ---------------------------------------------------------------- PDF

async def html_to_pdf(html_text: str, out: Path) -> int:
    """Print HTML to PDF with a headless Chromium browser (Chrome, Edge, Chromium or Brave); returns the page count."""
    chrome = osutil.find_chrome()
    if not chrome:
        raise RuntimeError("PDF export needs Google Chrome, Microsoft Edge or Chromium. Install one, or set JOBPORTAL_CHROME "
                           "to the browser's full path.")
    CACHE.mkdir(parents=True, exist_ok=True)
    out = Path(out).resolve()
    with tempfile.TemporaryDirectory(dir=CACHE, ignore_cleanup_errors=True) as tmp:
        src = Path(tmp) / "doc.html"
        src.write_text(html_text, encoding="utf-8")
        out.unlink(missing_ok=True)
        proc = await asyncio.create_subprocess_exec(
            chrome, "--headless=new", "--disable-gpu", "--no-sandbox", "--no-first-run", "--no-default-browser-check",
            "--disable-extensions", f"--user-data-dir={Path(tmp) / 'prof'}", "--no-pdf-header-footer", "--virtual-time-budget=1500",
            f"--print-to-pdf={out}", src.as_uri(),
            stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL, **osutil.popen_flags())
        try:
            await asyncio.wait_for(proc.wait(), timeout=90)
        except asyncio.TimeoutError:
            proc.kill()
            raise RuntimeError("PDF export timed out")
    if not out.exists():
        raise RuntimeError(f"the browser at {chrome} didn't write a PDF")
    return pdf_pages(out)


def pdf_pages(p: Path) -> int:
    return pdf_page_count(p)


async def resume_pdf(md: str, out: Path, theme: str = "classic", max_pages: int = 1) -> dict:
    """Render, shrinking slightly (down to 86%) until it fits the page limit (1, or 2 for federal resumes)."""
    out.parent.mkdir(parents=True, exist_ok=True)
    pages, fit = 0, 0.86
    for scale in (1.0, 0.97, 0.94, 0.91, 0.88, 0.86):
        pages = await html_to_pdf(resume_html(md, scale, theme), out)
        if pages <= max_pages:
            fit = scale
            break
    out.with_suffix(".fit").write_text(json.dumps({"scale": fit, "theme": theme, "pages": pages, "max_pages": max_pages}), encoding="utf-8")
    info = {"pages": pages, "scale": fit, "theme": theme, "max_pages": max_pages}
    if pages > max_pages:
        info["warning"] = f"Longer than {max_pages} page(s) even at 86%. Trim content."
    return info


def fit_info(pdf_path: Path) -> dict:
    try:
        raw = pdf_path.with_suffix(".fit").read_text(encoding="utf-8")
        return json.loads(raw) if raw.strip().startswith("{") else {"scale": float(raw)}
    except (OSError, ValueError):
        return {"scale": 1.0}


def fit_scale(pdf_path: Path) -> float:
    return float(fit_info(pdf_path).get("scale", 1.0))


async def cover_pdf(md: str, facts: dict, company: str, location: str, out: Path) -> dict:
    out.parent.mkdir(parents=True, exist_ok=True)
    pages = 0
    for scale in (1.0, 0.95, 0.9):
        pages = await html_to_pdf(cover_html(md, facts, company, location, scale), out)
        if pages <= 1:
            break
    return {"pages": pages}


# ---------------------------------------------------------------- DOCX

def _docx_runs(par, text: str):
    """Add inline-formatted runs (bold/italic/links) to a python-docx paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pos = 0
    tokens = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)|\*\*(.+?)\*\*|(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])|(https?://[^\s,)]+)")
    for m in tokens.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        if m.group(1) or m.group(5):
            label, url = (m.group(1), m.group(2)) if m.group(1) else (m.group(5), m.group(5))
            r_id = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), r_id)
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "1155CC")
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rpr.append(color)
            rpr.append(u)
            r.append(rpr)
            t = OxmlElement("w:t")
            t.text = label
            t.set(qn("xml:space"), "preserve")
            r.append(t)
            link.append(r)
            par._p.append(link)
        elif m.group(3):
            par.add_run(m.group(3)).bold = True
        elif m.group(4):
            par.add_run(m.group(4)).italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def _bottom_border(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def _new_doc(margins_in: float, font_pt: float):
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(margins_in)
        s.top_margin, s.bottom_margin = Inches(0.45 if margins_in < 0.8 else 0.9), Inches(0.4 if margins_in < 0.8 else 0.9)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(font_pt)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.0
    return doc


def resume_docx(md: str, out: Path, scale: float = 1.0, theme: str = "classic") -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt
    d = parse_resume(md)
    body_pt = {"compact": 9.6, "academic": 10.8}.get(theme, 10) * scale
    margin = 0.5 if scale >= 0.95 else 0.42
    doc = _new_doc(margin, body_pt)
    if theme == "academic":
        doc.styles["Normal"].font.name = "Times New Roman"
    if scale < 0.95:
        for sec in doc.sections:
            sec.top_margin = Inches(0.35)
    text_width = Inches(8.5 - 2 * margin)
    align = WD_ALIGN_PARAGRAPH.LEFT if theme == "modern" else WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(d["name"])
    r.bold, r.font.size = True, Pt({"modern": 19, "compact": 15, "academic": 18}.get(theme, 16.5) * scale)
    for c in d["contact"]:
        p = doc.add_paragraph()
        p.alignment = align
        _docx_runs(p, c)
        for run in p.runs:
            run.font.size = Pt((10 if theme in ("modern", "compact") else 11) * scale)
    head_pt = {"modern": 10.5, "compact": 10.5, "academic": 12}.get(theme, 11.5)
    for s in d["sections"]:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt((5 if theme == "compact" else 7) * scale)
        h.paragraph_format.space_after = Pt(2 * scale)
        rr = h.add_run(s["title"].upper() if theme == "modern" else s["title"])
        rr.bold, rr.font.size = True, Pt(head_pt * scale)
        if theme == "academic":
            rr.font.small_caps = True
        _bottom_border(h)
        for i, e in enumerate(s["entries"]):
            for kind, text in e:
                if kind == "bullet":
                    p = doc.add_paragraph(style="List Bullet")
                    p.paragraph_format.left_indent = Inches(0.3)
                    p.paragraph_format.space_after = Pt(0)
                    _docx_runs(p, text.replace(" || ", " | "))
                else:
                    p = doc.add_paragraph()
                    if " || " in text:  # right-aligned dates via a real right tab stop (ATS reads it as plain text)
                        left, right = text.split(" || ", 1)
                        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
                        _docx_runs(p, left)
                        p.add_run("\t")
                        _docx_runs(p, right)
                    else:
                        _docx_runs(p, text)
            if i < len(s["entries"]) - 1:
                p.paragraph_format.space_after = Pt((3.5 if theme == "compact" else 5) * scale)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def cover_docx(md: str, facts: dict, company: str, location: str, out: Path) -> None:
    from docx.shared import Pt
    doc = _new_doc(1.0, 11)
    p = doc.add_paragraph()
    r = p.add_run(facts.get("name", ""))
    r.bold, r.font.size = True, Pt(18)
    links = facts.get("links", {})
    p = doc.add_paragraph(" | ".join(x for x in [facts.get("email"), facts.get("phone"),
                                                  f"{facts.get('city', '')}, {facts.get('state', '')}".strip(", "),
                                                  links.get("site"), links.get("linkedin")] if x))
    _bottom_border(p)
    p.paragraph_format.space_after = Pt(16)
    doc.add_paragraph(long_date()).paragraph_format.space_after = Pt(10)
    if company:
        for l in ["Hiring Team", company] + ([location] if location else []):
            doc.add_paragraph(l)
        doc.paragraphs[-1].paragraph_format.space_after = Pt(14)
    for para in [x.strip() for x in re.split(r"\n\s*\n", (md or "").strip()) if x.strip()]:
        p = doc.add_paragraph()
        for i, l in enumerate(para.split("\n")):
            if i:
                p.add_run().add_break()
            _docx_runs(p, l)
        p.paragraph_format.space_after = Pt(11)
        p.paragraph_format.line_spacing = 1.15
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


# ---------------------------------------------------------------- generic markdown (briefs, prep) for previews

def markdown_html(md: str) -> str:
    out, in_ul = [], False
    for line in (md or "").split("\n"):
        st = line.strip()
        if re.match(r"^[-*•]\s+", st):
            if not in_ul:
                out.append("<ul>")
                in_ul = True
            item = re.sub(r"^[-*•]\s+", "", st)
            out.append(f"<li>{inline_html(item)}</li>")
            continue
        if in_ul:
            out.append("</ul>")
            in_ul = False
        if re.match(r"^\d+[.)]\s+", st):
            out.append(f"<p>{inline_html(st)}</p>")
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", st)
        if m:
            out.append(f"<h{len(m.group(1)) + 1}>{inline_html(m.group(2))}</h{len(m.group(1)) + 1}>")
        elif st:
            out.append(f"<p>{inline_html(st)}</p>")
    if in_ul:
        out.append("</ul>")
    return "\n".join(out)


# ---------------------------------------------------------------- generic documents (interview prep, intel, answers, emails)

DOC_CSS = """
@page { size: Letter; margin: 0; }
html, body { margin: 0; background: #fff; }
body { font-family: Arial, "Liberation Sans", Helvetica, sans-serif; color: #111; }
.page { width: 8.5in; padding: 0.75in 0.85in; font-size: 10.5pt; line-height: 1.45; }
.doc-title { font-size: 17pt; font-weight: 700; margin: 0 0 2pt; }
.doc-sub { color: #555; font-size: 9.5pt; margin-bottom: 14pt; border-bottom: 1px solid #000; padding-bottom: 6pt; }
h2 { font-size: 13pt; margin: 14pt 0 5pt; } h3 { font-size: 11.5pt; margin: 12pt 0 4pt; } h4, h5 { font-size: 10.5pt; margin: 10pt 0 3pt; }
p { margin: 0 0 7pt; } ul, ol { margin: 0 0 8pt; padding-left: 0.3in; } li { margin: 2pt 0; }
a { color: #1155cc; } code { font-family: monospace; font-size: 9.5pt; }
"""


def markdown_doc_html(md: str, title: str = "", subtitle: str = "") -> str:
    head = (f'<div class="doc-title">{html.escape(title)}</div>' if title else "") + \
           (f'<div class="doc-sub">{html.escape(subtitle)}</div>' if subtitle else "")
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title><style>{DOC_CSS}</style></head>'
            f'<body><div class="page">{head}{markdown_html(md)}</div></body></html>')


async def markdown_pdf(md: str, out: Path, title: str = "", subtitle: str = "") -> int:
    out.parent.mkdir(parents=True, exist_ok=True)
    return await html_to_pdf(markdown_doc_html(md, title, subtitle), out)


def markdown_docx(md: str, out: Path, title: str = "", subtitle: str = "") -> None:
    from docx.shared import Pt
    doc = _new_doc(0.85, 10.5)
    if title:
        r = doc.add_paragraph().add_run(title)
        r.bold, r.font.size = True, Pt(17)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        _bottom_border(p)
        p.paragraph_format.space_after = Pt(12)
    for line in (md or "").split("\n"):
        st = line.strip()
        if not st:
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", st)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2)))
            r.bold, r.font.size = True, Pt({1: 15, 2: 13, 3: 11.5}.get(len(m.group(1)), 10.5))
            continue
        if re.match(r"^[-*•]\s+", st):
            p = doc.add_paragraph(style="List Bullet")
            _docx_runs(p, re.sub(r"^[-*•]\s+", "", st))
        elif re.match(r"^\d+[.)]\s+", st):
            p = doc.add_paragraph(style="List Number")
            _docx_runs(p, re.sub(r"^\d+[.)]\s+", "", st))
        else:
            p = doc.add_paragraph()
            _docx_runs(p, st)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def eml(to: str, subject: str, body: str, sender: str = "") -> bytes:
    """An unsent email draft that opens as an editable draft in Thunderbird, Outlook and Apple Mail."""
    from email.message import EmailMessage
    msg = EmailMessage()
    if sender:
        msg["From"] = sender
    msg["To"] = to or ""
    msg["Subject"] = subject or ""
    msg["X-Unsent"] = "1"
    msg.set_content(body or "")
    return bytes(msg)
