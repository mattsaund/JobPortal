"""Text extraction for imported resumes. Pure Python for PDF, DOCX, ODT, RTF, Markdown and text, so it works the same on
Windows, macOS and Linux; old .doc / .pages / .wps files go through LibreOffice when it is installed."""
import io
import re
import subprocess
import tempfile
import zipfile
from pathlib import Path

from . import osutil

SUPPORTED = (".pdf", ".docx", ".odt", ".rtf", ".md", ".markdown", ".txt", ".doc", ".pages", ".wps")
ACCEPT = ",".join(SUPPORTED)
MIN_CHARS = 80


class ExtractError(ValueError):
    pass


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", "ignore")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", "ignore")


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x0c", "\n")
    text = re.sub(r"[ \t\xa0]{4,}", "   ", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            t = page.extract_text(extraction_mode="layout")  # keeps two-column resumes readable
        except Exception:  # noqa: BLE001 (layout mode is newer and stricter; plain mode reads anything)
            t = page.extract_text() or ""
        pages.append(t or "")
    return "\n\n".join(pages)


def pdf_page_count(path: Path) -> int:
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:  # noqa: BLE001
        return 0


def docx_text(data: bytes) -> str:
    """Paragraphs and tables in reading order, plus headers and footers (contact details often live there)."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for section in doc.sections:
        for part in (section.header, section.first_page_header):
            try:
                lines += [p.text for p in part.paragraphs if p.text.strip()]
            except Exception:  # noqa: BLE001
                pass
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                cells: list[str] = []
                for c in row.cells:
                    t = c.text.strip()
                    if t and t not in cells:  # merged cells repeat their text
                        cells.append(t)
                lines.append("   ".join(cells))
    for section in doc.sections:
        try:
            lines += [p.text for p in section.footer.paragraphs if p.text.strip()]
        except Exception:  # noqa: BLE001
            pass
    return "\n".join(lines)


def odt_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("content.xml").decode("utf-8", "ignore")
    xml = re.sub(r"<text:tab/>", "\t", xml)
    xml = re.sub(r"<text:line-break/>", "\n", xml)
    xml = re.sub(r"</text:(p|h)>", "\n", xml)
    xml = re.sub(r"<text:s(?: text:c=\"(\d+)\")?/>", lambda m: " " * int(m.group(1) or 1), xml)
    text = re.sub(r"<[^>]+>", "", xml)
    import html
    return html.unescape(text)


_RTF_DESTINATIONS = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer", "headerl", "headerr", "footerl",
                     "footerr", "listtable", "listoverridetable", "rsidtbl", "generator", "xmlnstbl", "themedata",
                     "colorschememapping", "latentstyles", "datastore", "object", "fldinst"}


def rtf_text(data: bytes) -> str:
    """A small RTF reader: keeps the text, paragraph breaks and tabs, drops formatting and embedded objects."""
    s = data.decode("latin-1", "ignore")
    out: list[str] = []
    stack: list[tuple[bool, int]] = []
    skip, uc, pending_skip = False, 1, 0
    i, n = 0, len(s)
    while i < n:
        ch = s[i]
        if ch == "{":
            stack.append((skip, uc))
            i += 1
        elif ch == "}":
            skip, uc = stack.pop() if stack else (False, 1)
            i += 1
        elif ch == "\\":
            m = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\(.)", s[i:])
            if not m:
                i += 1
                continue
            i += m.end()
            word, arg, hexc, sym = m.group(1), m.group(2), m.group(3), m.group(4)
            if hexc:
                if pending_skip:
                    pending_skip -= 1
                elif not skip:
                    out.append(bytes([int(hexc, 16)]).decode("cp1252", "ignore"))
            elif sym:
                if sym == "*":
                    skip = True
                elif not skip and sym in "\\{}":
                    out.append(sym)
                elif not skip and sym == "~":
                    out.append(" ")
            elif word in _RTF_DESTINATIONS:
                skip = True
            elif skip:
                continue
            elif word in ("par", "line", "sect", "page", "row"):
                out.append("\n")
            elif word == "tab" or word == "cell":
                out.append("\t")
            elif word == "uc":
                uc = int(arg or 1)
            elif word == "u":
                code = int(arg or 0)
                out.append(chr(code + 65536 if code < 0 else code))
                pending_skip = uc
        else:
            if pending_skip and ch not in "\r\n":
                pending_skip -= 1
            elif not skip and ch not in "\r\n":
                out.append(ch)
            i += 1
    return "".join(out)


def _soffice_text(name: str, data: bytes) -> str:
    soffice = osutil.find_soffice()
    if not soffice:
        raise ExtractError(f"{Path(name).suffix} files need LibreOffice to read. Save the resume as PDF or DOCX and import that instead.")
    from .config import CACHE
    CACHE.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=CACHE, ignore_cleanup_errors=True) as tmp:
        src = Path(tmp) / f"resume{Path(name).suffix.lower()}"
        src.write_bytes(data)
        profile = (Path(tmp) / "lo-profile").as_uri()
        subprocess.run([soffice, "--headless", f"-env:UserInstallation={profile}", "--convert-to", "txt:Text", "--outdir", tmp, str(src)],
                       capture_output=True, timeout=180, **osutil.popen_flags())
        out = Path(tmp) / "resume.txt"
        return out.read_text(encoding="utf-8", errors="ignore") if out.exists() else ""


def extract_text(name: str, data: bytes) -> str:
    """Plain text of one resume file. Raises ExtractError with a message meant for the user."""
    suf = Path(name or "").suffix.lower()
    try:
        if suf in (".md", ".markdown", ".txt", ""):
            text = _decode(data)
        elif suf == ".pdf":
            text = pdf_text(data)
        elif suf == ".docx":
            text = docx_text(data)
        elif suf == ".odt":
            text = odt_text(data)
        elif suf == ".rtf":
            text = rtf_text(data)
        elif suf in (".doc", ".pages", ".wps"):
            text = _soffice_text(name, data)
        else:
            raise ExtractError(f"{suf} files aren't supported. Use PDF, DOCX, ODT, RTF, Markdown or plain text.")
    except ExtractError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError, OSError, subprocess.SubprocessError) as e:
        raise ExtractError(f"couldn't read {name}: {e}") from e
    except Exception as e:  # noqa: BLE001 (third-party parsers raise their own error types)
        raise ExtractError(f"couldn't read {name}: {type(e).__name__}: {e}") from e
    text = _tidy(text)
    if len(text) < MIN_CHARS:
        raise ExtractError("couldn't find any text in that file. A scanned (image) PDF can't be read; export a text PDF or DOCX instead.")
    return text
