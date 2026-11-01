"""Resume, cover letter and document rendering: Markdown in, HTML, PDF and Word (DOCX) files out.

Who calls it: documents/exports.py builds every download with it, search/pipeline.py renders each new application package
and measures how many pages a resume prints to, the web routes show live previews in the editor, and
documents/resumecheck.py reuses parse_resume().

How a resume is made:
  1. parse_resume() reads the resume Markdown (the format is RESUME_FORMAT in ai/prompts.py) into a small structure: the
     name, the contact lines, and sections made of entries (one job, school or project each).
  2. resume_html() turns that structure into one HTML page, styled by the chosen layout (a theme in RESUME_THEMES).
  3. resume_pdf() prints the page to PDF with a headless Chromium browser (html_to_pdf). If the resume runs past its page
     limit, it prints again with slightly smaller text, down to 86%, until it fits. That is the fit loop.
  4. resume_docx() builds the same resume in Word with python-docx, at the text size the PDF settled on, so both files look
     alike and both fit the page.

Cover letters (cover_html, cover_pdf, cover_docx) and other documents such as interview prep and the company brief
(markdown_html, markdown_pdf, markdown_docx) follow the same path with simpler layouts. Every page is US Letter.

Safety: the Markdown comes from the AI and from job posts on the open web, so all of its text is HTML-escaped and only web
and mail links are kept (safe_href). python-docx is imported inside the DOCX functions, so it loads only when a DOCX is made."""
import asyncio
import html
import json
import os
import re
import tempfile
import uuid
from datetime import date
from email.message import EmailMessage
from itertools import groupby
from pathlib import Path

from jobportal.core import osutil
from jobportal.core.config import CACHE
from jobportal.documents.extract import pdf_page_count


def long_date(d: date | None = None) -> str:
    """A date like "October 5, 2026" (today by default). Built by hand because strftime's %-d (day without a leading zero)
    doesn't exist on Windows."""
    d = d or date.today()
    return f"{d:%B} {d.day}, {d.year}"


# ---------------------------------------------------------------- inline formatting (shared by HTML and DOCX)

# The inline Markdown a line may use: RESUME_FORMAT allows **bold**, *italic* and [text](url), and bare web addresses become
# links too. One pattern finds them all, left to right. Where several could start at the same spot, the first one listed
# wins, so a link beats bold, and *** (bold and italic) beats **.
_INLINE_RE = re.compile(
    r"\[(?P<label>[^\]]+)\]\((?P<href>[^)\s]+)\)"  # [text](url)
    r"|\*\*\*(?P<bold_italic>.+?)\*\*\*"  # ***bold and italic***
    r"|\*\*(?P<bold>.+?)\*\*"  # **bold**
    r"|(?<![*\w])\*(?!\s)(?P<italic>.+?)(?<!\s)\*(?![*\w])"  # *italic*, but not the stars in 2*3*4 or "a * b"
    # A bare address. It can't follow a letter, digit, slash, quote or "=" (then it is part of something else). It stops at
    # spaces, commas, ")", "*" and the characters that could break out of HTML, and it never ends on punctuation, so the
    # period in "see https://example.com." stays out of the link.
    r"|(?<![\w/\"'=])(?P<url>https?://[^\s,)<>\"*]*[^\s,)<>\"*.!?:;'])"
)


def safe_href(url: str) -> str:
    """The link itself when it is a web or mail address, else "#". AI output and job posts come from the open web, and a
    javascript: or file: link must never reach a document."""
    url = (url or "").strip()
    return url if re.match(r"^(https?://|mailto:)", url, re.I) else "#"


def _inline_pieces(text: str, bold: bool = False, italic: bool = False) -> list[tuple[str, str, bool, bool]]:
    """Split one line into (text, href, bold, italic) pieces, where href is "" for plain text and "#" for a link that
    safe_href refused. Bold and italic text can hold links, so its inside is split again with the style switched on.
    inline_html and _docx_runs both render these pieces, so the PDF and the DOCX always read a line the same way."""
    pieces, pos = [], 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            pieces.append((text[pos:m.start()], "", bold, italic))
        if m["label"] is not None:
            pieces.append((m["label"], safe_href(m["href"]), bold, italic))
        elif m["url"] is not None:
            pieces.append((m["url"], m["url"], bold, italic))
        elif m["bold_italic"] is not None:
            pieces += _inline_pieces(m["bold_italic"], True, True)
        elif m["bold"] is not None:
            pieces += _inline_pieces(m["bold"], True, italic)
        else:
            pieces += _inline_pieces(m["italic"], bold, True)
        pos = m.end()
    if pos < len(text):
        pieces.append((text[pos:], "", bold, italic))
    return pieces


def inline_html(text: str) -> str:
    """One line of Markdown as safe HTML: text is escaped, **bold** and *italic* become <b> and <i>, links become <a>."""
    out = []
    for piece, href, bold, italic in _inline_pieces(text):
        part = f'<a href="{html.escape(href)}">{html.escape(piece)}</a>' if href else html.escape(piece, quote=False)
        if italic:
            part = f"<i>{part}</i>"
        if bold:
            part = f"<b>{part}</b>"
        out.append(part)
    return "".join(out)


# ---------------------------------------------------------------- resume parsing

_BULLET_RE = re.compile(r"^[-*\u2022]\s+")  # "- ", "* " or "• " (a bullet character) starts a bullet line


def parse_resume(md: str) -> dict:
    """Read resume Markdown into {"name": str, "contact": [line], "sections": [{"title": str, "entries": [entry]}]}.

    The rules follow RESUME_FORMAT in ai/prompts.py. The first "# " line is the name. Lines before the first "## " heading
    are contact lines. "## Title" starts a section. Inside a section a blank line ends the current entry (one job, school or
    project), and an entry is a list of ("bullet", text) for "- " lines and ("line", text) for every other line."""
    resume = {"name": "", "contact": [], "sections": []}
    section = None
    entry: list | None = None
    for raw in (md or "").replace("\r", "").split("\n"):
        line = raw.strip()
        if line.startswith("# ") and not resume["name"]:
            resume["name"] = line[2:].strip()
        elif line.startswith("## "):
            section = {"title": line[3:].strip(), "entries": []}
            resume["sections"].append(section)
            entry = None
        elif not line:
            entry = None
        elif section is None:
            resume["contact"].append(line)
        else:
            if entry is None:
                entry = []
                section["entries"].append(entry)
            if _BULLET_RE.match(line):
                entry.append(("bullet", _BULLET_RE.sub("", line)))
            else:
                entry.append(("line", line.lstrip("#").strip()))  # a stray "### Org" heading reads as a plain line
    return resume


# ---------------------------------------------------------------- resume HTML

# How the resume page is built:
#   .page is one sheet of US Letter, and its padding is the page margin, so the editor's preview shows the margins too.
#   @page { margin: 0 } hands the whole sheet to .page when printing.
#   Every size is multiplied by --s, the text scale the fit loop picks (1 is 100%). --mt and --mx are the top and side
#   margins. resume_html sets all three on .page, and the editor preview (web/static/js/pages/lead.js) changes them the same
#   way to mirror the fit loop, so keep FIT_SCALES and _resume_margins in step with it.
#   box-decoration-break: clone repeats the padding on every printed page. Without it a two-page resume would run to the
#   bottom edge of page 1 and start at the very top of page 2.
#   A "left || right" line is a flex row with the right part (usually dates) pushed flush right.
RESUME_BASE_CSS = """
@page { size: Letter; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; }
body { font-family: Arial, "Liberation Sans", Helvetica, sans-serif; color: #000; -webkit-print-color-adjust: exact; }
.page { width: 8.5in; min-height: 11in; padding: var(--mt, 0.45in) var(--mx, 0.5in) 0.4in; font-size: calc(10pt * var(--s, 1)); line-height: 1.17;
        box-decoration-break: clone; }
h1 { font-size: calc(16.5pt * var(--s, 1)); text-align: center; margin: 0 0 3pt; font-weight: 700; }
.contact { text-align: center; font-size: calc(11pt * var(--s, 1)); margin-bottom: 2pt; }
a { color: #1155cc; text-decoration: underline; }
h2 { font-size: calc(11.5pt * var(--s, 1)); font-weight: 700; margin: calc(7pt * var(--s, 1)) 0 calc(2pt * var(--s, 1));
     border-bottom: 1px solid #000; padding-bottom: 1pt; }
.entry { margin: 0 0 calc(5pt * var(--s, 1)); }
.entry:last-child { margin-bottom: 0; }
.line { margin: 0; }
.line.split { display: flex; justify-content: space-between; gap: 12pt; }
.line.split .r { white-space: nowrap; }
ul { margin: 0; padding-left: 0.26in; }
li { margin: 0; padding-left: 0.04in; }
"""

# The layouts a user can pick. Every one is single-column and text only, so applicant tracking systems (ATS) read it
# cleanly; they differ only in fonts, sizes and header alignment. Each is CSS added after RESUME_BASE_CSS, so it lists only
# what it changes. resume_docx copies the same choices into Word from _DOCX_THEMES, so change a layout in both places.
#   classic   the base styles: Arial, centered name and contact line, ruled section headings
#   modern    left-aligned header, a bigger name, uppercase headings with wide letter spacing, black links with no underline
#   compact   smaller text and tighter spacing, for a full one-page resume
#   academic  Times New Roman with small-caps headings, the usual look for research and CV-style resumes
RESUME_THEMES = {
    "classic": "",
    "modern": """
h1 { text-align: left; font-size: calc(19pt * var(--s, 1)); margin-bottom: 2pt; letter-spacing: .2pt; }
.contact { text-align: left; font-size: calc(10pt * var(--s, 1)); color: #222; }
h2 { font-size: calc(10.5pt * var(--s, 1)); text-transform: uppercase; letter-spacing: .9pt; margin-top: calc(8pt * var(--s, 1)); }
a { color: #000; text-decoration: none; }
""",
    "compact": """
.page { font-size: calc(9.6pt * var(--s, 1)); line-height: 1.12; }
h1 { font-size: calc(15pt * var(--s, 1)); margin-bottom: 1pt; }
.contact { font-size: calc(10pt * var(--s, 1)); }
h2 { font-size: calc(10.5pt * var(--s, 1)); margin: calc(5pt * var(--s, 1)) 0 calc(1.5pt * var(--s, 1)); }
.entry { margin-bottom: calc(3.5pt * var(--s, 1)); }
""",
    "academic": """
body { font-family: "Times New Roman", "Liberation Serif", Times, serif; }
.page { font-size: calc(10.8pt * var(--s, 1)); line-height: 1.15; }
h1 { font-size: calc(18pt * var(--s, 1)); font-weight: 700; }
h2 { font-size: calc(12pt * var(--s, 1)); font-variant: small-caps; letter-spacing: .4pt; }
""",
}
THEMES = list(RESUME_THEMES)
# A resume's layout and page limit when nothing chose them: a package's resume_style, a saved resume's style.
DEFAULT_STYLE = {"theme": "classic", "max_pages": 1}

# The text scales the fit loop tries, largest first. lead.js mirrors this list for the editor preview.
FIT_SCALES = (1.0, 0.97, 0.94, 0.91, 0.88, 0.86)
RESUME_BOTTOM_MARGIN = 0.4  # inches, at every scale


def _resume_margins(scale: float) -> tuple[float, float]:
    """Top and side margins in inches for a text scale. Below 95% the margins tighten as well, which wins room for a long
    resume before the text has to shrink any further."""
    return (0.45, 0.47) if scale >= 0.95 else (0.35, 0.42)


def _line_html(text: str) -> str:
    """One resume line. "left || right" puts the right part (usually dates) flush right on the same line."""
    if " || " in text:
        left, right = text.split(" || ", 1)
        return f'<div class="line split"><span>{inline_html(left)}</span><span class="r">{inline_html(right)}</span></div>'
    return f'<div class="line">{inline_html(text)}</div>'


def _entry_html(entry: list[tuple[str, str]]) -> str:
    """One entry (a job, school or project). Consecutive bullets share one <ul>; every other line is its own row."""
    parts = ['<div class="entry">']
    for is_bullet, items in groupby(entry, key=lambda item: item[0] == "bullet"):
        if is_bullet:  # a bullet has no right-hand column, so a stray " || " reads as a plain separator
            parts += ["<ul>", *(f"<li>{inline_html(text.replace(' || ', ' | '))}</li>" for _, text in items), "</ul>"]
        else:
            parts += [_line_html(text) for _, text in items]
    parts.append("</div>")
    return "".join(parts)


def resume_html(md: str, scale: float = 1.0, theme: str = "classic") -> str:
    """A resume as one HTML page in a layout from THEMES, with every size multiplied by scale. This page is exactly what
    resume_pdf prints and what the editor shows as its live preview. An unknown theme falls back to classic."""
    resume = parse_resume(md)
    parts = [f"<h1>{inline_html(resume['name'])}</h1>"]
    parts += [f'<div class="contact">{inline_html(line)}</div>' for line in resume["contact"]]
    for section in resume["sections"]:
        parts.append(f"<h2>{inline_html(section['title'])}</h2>")
        parts += [_entry_html(entry) for entry in section["entries"]]
    top, side = _resume_margins(scale)
    css = RESUME_BASE_CSS + RESUME_THEMES.get(theme, "")
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(resume["name"])} Resume</title>'
            f'<style>{css}</style></head><body><div class="page" style="--s:{scale};--mt:{top}in;--mx:{side}in">'
            + "\n".join(parts) + "</div></body></html>")


# ---------------------------------------------------------------- cover letter HTML

# A business letter: name and contact line over a rule, then the date, the recipient and the paragraphs. Only the body
# text follows the scale (--s); cover_pdf shrinks it a little when a letter runs long.
COVER_CSS = """
@page { size: Letter; margin: 0; }
html, body { margin: 0; padding: 0; background: #fff; }
body { font-family: Arial, "Liberation Sans", Helvetica, sans-serif; color: #000; }
.page { width: 8.5in; min-height: 11in; padding: 0.9in 1in; font-size: calc(11pt * var(--s, 1)); line-height: 1.42; box-decoration-break: clone; }
.hd h1 { font-size: 18pt; margin: 0 0 2pt; }
.hd .contact { font-size: 10pt; color: #222; }
.hd { border-bottom: 1px solid #000; padding-bottom: 8pt; margin-bottom: 18pt; }
.meta { margin-bottom: 14pt; }
p { margin: 0 0 11pt; }
a { color: #1155cc; }
"""


def _cover_contact(facts: dict) -> str:
    """The contact line under the name on a cover letter: email | phone | City, ST | website | LinkedIn. The web addresses
    lose their "https://" (and LinkedIn its "www.") to keep the line short."""
    links = facts.get("links") or {}
    place = ", ".join(x for x in (facts.get("city"), facts.get("state")) if x)
    site = (links.get("site") or "").replace("https://", "")
    linkedin = (links.get("linkedin") or "").replace("https://www.", "")
    return " | ".join(x for x in (facts.get("email"), facts.get("phone"), place, site, linkedin) if x)


def _recipient_lines(company: str, location: str) -> list[str]:
    """Who the letter goes to, under the date. Nothing when the company is unknown."""
    if not company:
        return []
    return ["Hiring Team", company] + ([location] if location else [])


def _letter_paragraphs(md: str) -> list[list[str]]:
    """A letter's paragraphs (split at blank lines), each as its list of lines."""
    return [para.strip().split("\n") for para in re.split(r"\n\s*\n", (md or "").strip()) if para.strip()]


def cover_html(md: str, facts: dict, company: str = "", location: str = "", scale: float = 1.0) -> str:
    """A cover letter as one HTML page. facts is the candidate profile's facts (name, email, phone, city, state, links).
    The date is today's. Lines inside a paragraph stay separate lines, as they were typed."""
    body = "".join(f"<p>{'<br>'.join(inline_html(line) for line in lines)}</p>" for lines in _letter_paragraphs(md))
    meta = long_date()
    recipient = _recipient_lines(company, location)
    if recipient:
        meta += "<br><br>" + "<br>".join(html.escape(line) for line in recipient)
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>Cover Letter</title><style>{COVER_CSS}</style></head>'
            f'<body><div class="page" style="--s:{scale}"><div class="hd"><h1>{html.escape(facts.get("name", ""))}</h1>'
            f'<div class="contact">{inline_html(_cover_contact(facts))}</div></div><div class="meta">{meta}</div>{body}</div></body></html>')


# ---------------------------------------------------------------- PDF

PDF_TIMEOUT = 90  # seconds for the browser to print one page set

# Switches for printing with a headless Chromium browser.
_PRINT_FLAGS = (
    "--headless=new",  # no window
    "--disable-gpu",  # printing needs no graphics card, and some drivers crash in headless mode
    # Chrome's sandbox can't start in some setups: as root, in containers, or where Linux blocks user namespaces for programs
    # outside the system's own (recent Ubuntu does). The page is JobPortal's own escaped HTML with no scripts.
    "--no-sandbox",
    "--no-first-run", "--no-default-browser-check", "--disable-extensions",  # a bare browser: no welcome pages or add-ons
    "--no-pdf-header-footer",  # no date, title or address printed in the margins
    "--virtual-time-budget=1500",  # let the page finish its layout (fonts included) before printing
)


class PDFError(RuntimeError):
    """Printing a PDF failed. The message is written for the user: the web routes show it as it is (a 503)."""


async def html_to_pdf(html_text: str, out: Path) -> int:
    """Print an HTML page to the PDF file out and return its page count.

    The printer is osutil.find_pdf_engine(): JOBPORTAL_CHROME when set, else the engine bundled with the app, else an
    installed Chrome, Edge, Chromium or Brave. It runs as a headless subprocess. Raises PDFError when there is no such
    browser, when it takes longer than PDF_TIMEOUT, or when it doesn't write a readable PDF."""
    browser = osutil.find_pdf_engine()
    if not browser:
        raise PDFError("PDF export needs Google Chrome, Microsoft Edge or Chromium. Install one, or set JOBPORTAL_CHROME "
                       "to the browser's full path.")
    CACHE.mkdir(parents=True, exist_ok=True)
    out = Path(out).resolve()  # the browser runs in its own process, so give it a full path
    out.unlink(missing_ok=True)  # an old file must never pass for a new one
    with tempfile.TemporaryDirectory(dir=CACHE, ignore_cleanup_errors=True) as tmp:
        page = Path(tmp) / "doc.html"
        page.write_text(html_text, encoding="utf-8")
        # A throwaway profile lets printing work while the user's own browser is open. With the default profile, Chrome would
        # hand the job to the running browser and exit without printing anything.
        profile = Path(tmp) / "profile"
        proc = await asyncio.create_subprocess_exec(
            browser, *_PRINT_FLAGS, f"--user-data-dir={profile}", f"--print-to-pdf={out}", page.as_uri(),
            stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL, **osutil.popen_flags())
        try:
            await asyncio.wait_for(proc.wait(), timeout=PDF_TIMEOUT)
        except asyncio.TimeoutError:
            proc.kill()
            await proc.wait()
            raise PDFError("PDF export timed out") from None
        except asyncio.CancelledError:
            proc.kill()  # Stop, or JobPortal shutting down: don't leave the browser printing on its own
            raise
    pages = pdf_page_count(out) if out.exists() else 0
    if not pages:
        raise PDFError(f"the browser at {browser} didn't write a PDF")
    return pages


async def _print_to_fit(make_html, scales: tuple[float, ...], max_pages: int, out: Path, before_publish=None) -> tuple[int, float]:
    """The fit loop: print make_html(scale) at each scale in turn, stopping at the first that fits max_pages (the last one
    is kept when none does), then move the result into place as out. Returns (pages, scale).

    Every try goes to a draft file beside out, and the finished PDF replaces out in one step. Downloads reuse a rendered
    file whenever it exists (documents/exports.py), so out must never hold a try that overflows, and two renders of the
    same file must not delete each other's work. before_publish(pages, scale) runs just before the move (resume_pdf writes
    the .fit file there, so a PDF never appears without the scale it was printed at)."""
    out.parent.mkdir(parents=True, exist_ok=True)
    draft = out.with_name(f".{out.stem}-{uuid.uuid4().hex[:8]}.draft{out.suffix}")
    try:
        for scale in scales:
            pages = await html_to_pdf(make_html(scale), draft)
            if pages <= max_pages:
                break
        if before_publish:
            before_publish(pages, scale)
        os.replace(draft, out)
    finally:
        draft.unlink(missing_ok=True)  # only still there when printing failed or was canceled
    return pages, scale


async def resume_pdf(md: str, out: Path, theme: str = "classic", max_pages: int = 1) -> dict:
    """Print a resume to the PDF file out, shrinking it until it fits max_pages (1, or 2 for federal resumes).

    The fit loop (_print_to_fit) prints at each scale in FIT_SCALES and stops at the first one that fits. When even 86% is
    too long, the 86% version is kept (resumecheck reports the overflow). Returns {"pages", "scale", "theme", "max_pages"}
    and saves the same facts next to the PDF (resume.pdf gets resume.fit), so resume_docx can later use the same scale
    (fit_scale)."""
    def save_fit(pages: int, scale: float) -> None:
        info.update(pages=pages, scale=scale)
        out.with_suffix(".fit").write_text(json.dumps(info), encoding="utf-8")

    info = {"pages": 0, "scale": 1.0, "theme": theme, "max_pages": max_pages}
    await _print_to_fit(lambda scale: resume_html(md, scale, theme), FIT_SCALES, max_pages, out, save_fit)
    return info


def fit_scale(pdf_path: Path) -> float:
    """The text scale resume_pdf settled on for this PDF, read from the .fit file beside it. 1.0 when there is none."""
    try:
        return float(json.loads(pdf_path.with_suffix(".fit").read_text(encoding="utf-8"))["scale"])
    except (OSError, ValueError, KeyError, TypeError):
        return 1.0


COVER_SCALES = (1.0, 0.95, 0.9)  # a cover letter shrinks less than a resume before it may take a second page


async def cover_pdf(md: str, facts: dict, company: str, location: str, out: Path) -> dict:
    """Print a cover letter to the PDF file out, shrinking the text a little when it runs past one page (the same fit loop
    as a resume, _print_to_fit). Returns {"pages"}."""
    pages, _ = await _print_to_fit(lambda scale: cover_html(md, facts, company, location, scale), COVER_SCALES, 1, out)
    return {"pages": pages}


# ---------------------------------------------------------------- DOCX building blocks

# Characters XML (and so a Word file) can't hold. python-docx refuses a whole document over one of them, and they do turn
# up: text pasted from Word carries \x0b for a manual line break, and PDF text can carry stray control codes.
_NOT_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xml_safe(value):
    """`value` with characters Word can't store turned into spaces, so one stray control character can't make every Word
    download fail. Strings inside a dict (the facts, and their links) are cleaned too; anything else comes back as it is."""
    if isinstance(value, str):
        return _NOT_XML_RE.sub(" ", value)
    if isinstance(value, dict):
        return {key: _xml_safe(item) for key, item in value.items()}
    return value


def _save(doc, out: Path) -> None:
    """Save a python-docx document at out, making its folder first."""
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def _new_docx(top: float, side: float, bottom: float, font_pt: float, font: str = "Arial"):
    """A blank Word document on US Letter with these margins (inches) and body text. Paragraphs get no spacing of their own;
    the builders add space exactly where the PDF has it."""
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(side)
        section.top_margin, section.bottom_margin = Inches(top), Inches(bottom)
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(font_pt)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0  # Word's single spacing is close to the CSS line heights
    return doc


def _add_link(paragraph, text: str, url: str, run_style: dict, color: str, underline: bool) -> None:
    """Add a clickable link. python-docx can read links but not add them, so this adds a normal run, then wraps its XML in
    Word's <w:hyperlink> element. The address itself goes in the document's list of relationships, which r:id points to."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    run = _add_run(paragraph, text, **run_style)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.underline = underline
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run.element.addprevious(link)  # put the link where the run is, then move the run inside it
    link.append(run.element)


def _add_run(paragraph, text: str, bold: bool, italic: bool, size_pt: float | None):
    """Add a run of text. Styles that are off stay unset, so the paragraph's own style shows through."""
    from docx.shared import Pt
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if size_pt:
        run.font.size = Pt(size_pt)
    return run


def _docx_runs(paragraph, text: str, size_pt: float | None = None, link_color: str = "1155CC", underline_links: bool = True) -> None:
    """Add one line of Markdown to a Word paragraph: plain, bold and italic runs plus real links, read by the same
    _inline_pieces as the HTML. size_pt sets a size on every run (links included). A refused link becomes plain text."""
    for piece, href, bold, italic in _inline_pieces(text):
        run_style = {"bold": bold, "italic": italic, "size_pt": size_pt}
        if href and href != "#":
            _add_link(paragraph, piece, href, run_style, link_color, underline_links)
        else:
            _add_run(paragraph, piece, **run_style)


def _bottom_border(paragraph) -> None:
    """Rule a thin line under a paragraph, Word's version of the CSS border-bottom under headings."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    border = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):  # sz is in eighths of a point
        border.set(qn(key), value)
    borders = OxmlElement("w:pBdr")
    borders.append(border)
    paragraph._p.get_or_add_pPr().append(borders)


# ---------------------------------------------------------------- resume DOCX

# Word versions of the choices in RESUME_BASE_CSS and RESUME_THEMES. Sizes and spacing are points at 100%; resume_docx
# multiplies them by the fit scale like the CSS does. Change a layout here and in its CSS together. Not copied: line heights
# (Word's single spacing is close) and letter spacing.
_DOCX_BASE = {"font": "Arial", "align": "center", "body": 10, "name": 16.5, "contact": 11, "heading": 11.5, "heading_before": 7,
              "heading_after": 2, "entry_gap": 5, "uppercase": False, "small_caps": False, "link_color": "1155CC", "link_underline": True}
_DOCX_THEMES = {
    "classic": {},
    "modern": {"align": "left", "name": 19, "contact": 10, "heading": 10.5, "heading_before": 8, "uppercase": True,
               "link_color": "000000", "link_underline": False},
    "compact": {"body": 9.6, "name": 15, "contact": 10, "heading": 10.5, "heading_before": 5, "heading_after": 1.5, "entry_gap": 3.5},
    "academic": {"font": "Times New Roman", "body": 10.8, "name": 18, "heading": 12, "small_caps": True},
}
BULLET_INDENT = 0.3  # inches: the CSS list indent (0.26in) plus the item padding (0.04in)


def resume_docx(md: str, out: Path, scale: float = 1.0, theme: str = "classic") -> None:
    """Write a resume as a Word file at out that matches resume_html in the same layout. Pass the scale the PDF settled on
    (fit_scale) so the DOCX keeps to the same number of pages.

    How the PDF maps to Word: the name, contact lines and section headings become paragraphs with the theme's sizes; a
    heading's rule is a paragraph border; bullets use Word's List Bullet style; and a "left || right" line uses a right tab
    stop at the right margin, so the dates sit flush right yet stay plain text an ATS reads in order. Word stores font sizes
    in half points and python-docx rounds down (9.4pt becomes 9pt), which keeps the DOCX text no bigger than the PDF's."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt
    style = {**_DOCX_BASE, **_DOCX_THEMES.get(theme, {})}
    links = {"link_color": style["link_color"], "underline_links": style["link_underline"]}
    resume = parse_resume(_xml_safe(md))
    top, side = _resume_margins(scale)
    doc = _new_docx(top, side, RESUME_BOTTOM_MARGIN, style["body"] * scale, style["font"])
    align = WD_ALIGN_PARAGRAPH.LEFT if style["align"] == "left" else WD_ALIGN_PARAGRAPH.CENTER

    name = doc.add_paragraph()
    name.alignment = align
    run = name.add_run(resume["name"])
    run.bold, run.font.size = True, Pt(style["name"] * scale)
    for line in resume["contact"]:
        contact = doc.add_paragraph()
        contact.alignment = align
        _docx_runs(contact, line, size_pt=style["contact"] * scale, **links)

    right_edge = Inches(8.5 - 2 * side)  # where the text area ends
    for section in resume["sections"]:
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(style["heading_before"] * scale)
        heading.paragraph_format.space_after = Pt(style["heading_after"] * scale)
        run = heading.add_run(section["title"].upper() if style["uppercase"] else section["title"])
        run.bold, run.font.size = True, Pt(style["heading"] * scale)
        if style["small_caps"]:
            run.font.small_caps = True
        _bottom_border(heading)
        for number, entry in enumerate(section["entries"]):
            for kind, text in entry:
                if kind == "bullet":
                    p = doc.add_paragraph(style="List Bullet")
                    p.paragraph_format.left_indent = Inches(BULLET_INDENT)
                    p.paragraph_format.space_after = Pt(0)
                    _docx_runs(p, text.replace(" || ", " | "), **links)
                elif " || " in text:
                    left, right = text.split(" || ", 1)
                    p = doc.add_paragraph()
                    p.paragraph_format.tab_stops.add_tab_stop(right_edge, WD_TAB_ALIGNMENT.RIGHT)
                    _docx_runs(p, left, **links)
                    p.add_run("\t")
                    _docx_runs(p, right, **links)
                else:
                    p = doc.add_paragraph()
                    _docx_runs(p, text, **links)
            if number < len(section["entries"]) - 1:  # the gap between entries goes after an entry's last line
                p.paragraph_format.space_after = Pt(style["entry_gap"] * scale)
    _save(doc, out)


# ---------------------------------------------------------------- cover letter DOCX

def cover_docx(md: str, facts: dict, company: str, location: str, out: Path) -> None:
    """Write a cover letter as a Word file at out, laid out like cover_html: name, contact line over a rule, date,
    recipient, then the paragraphs with their line breaks kept."""
    from docx.shared import Pt
    md, facts, company, location = _xml_safe(md), _xml_safe(facts), _xml_safe(company), _xml_safe(location)
    doc = _new_docx(top=0.9, side=1.0, bottom=0.9, font_pt=11)
    run = doc.add_paragraph().add_run(facts.get("name", ""))
    run.bold, run.font.size = True, Pt(18)
    contact = doc.add_paragraph()
    _docx_runs(contact, _cover_contact(facts), size_pt=10)
    _bottom_border(contact)
    contact.paragraph_format.space_after = Pt(16)
    doc.add_paragraph(long_date()).paragraph_format.space_after = Pt(10)
    recipient = _recipient_lines(company, location)
    for line in recipient:
        doc.add_paragraph(line)
    if recipient:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(14)
    for lines in _letter_paragraphs(md):
        p = doc.add_paragraph()
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break()  # a line break inside the paragraph, like <br> in the PDF
            _docx_runs(p, line)
        p.paragraph_format.space_after = Pt(11)
        p.paragraph_format.line_spacing = 1.15
    _save(doc, out)


# ---------------------------------------------------------------- generic documents (interview prep, company brief, answers, emails)

_HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)")


def _markdown_lines(md: str) -> list[tuple[str, str, int]]:
    """Sort each line of a generic document into (kind, text, heading level): kind is "heading", "bullet", "text" or
    "blank", and the level is 0 except for headings. Numbered lines ("1. ...") stay text with their number, so they print
    exactly as written. Word's own numbering would count every list in the document as one."""
    lines = []
    for raw in (md or "").split("\n"):
        line = raw.strip()
        heading = _HEADING_RE.match(line)
        if _BULLET_RE.match(line):
            lines.append(("bullet", _BULLET_RE.sub("", line), 0))
        elif heading:
            lines.append(("heading", heading.group(2), len(heading.group(1))))
        else:
            lines.append(("text" if line else "blank", line, 0))
    return lines


def markdown_html(md: str) -> str:
    """Generic Markdown as HTML: headings, bullet lists, one paragraph per line, and inline formatting. "#" becomes <h2>,
    because a document's title (markdown_doc_html) is the biggest text on its page. Used for previews and document PDFs."""
    out = []
    for is_bullet, lines in groupby(_markdown_lines(md), key=lambda line: line[0] == "bullet"):
        if is_bullet:  # consecutive bullets form one list; a blank or any other line ends it
            out += ["<ul>", *(f"<li>{inline_html(text)}</li>" for _, text, _ in lines), "</ul>"]
            continue
        for kind, text, level in lines:
            if kind == "heading":
                out.append(f"<h{level + 1}>{inline_html(text)}</h{level + 1}>")
            elif kind == "text":
                out.append(f"<p>{inline_html(text)}</p>")
    return "\n".join(out)


# A plain document page: optional title and subtitle over a rule, then the Markdown. Its margins are .page's padding, as on
# the resume, and markdown_docx copies them.
DOC_CSS = """
@page { size: Letter; margin: 0; }
html, body { margin: 0; background: #fff; }
body { font-family: Arial, "Liberation Sans", Helvetica, sans-serif; color: #111; }
.page { width: 8.5in; padding: 0.75in 0.85in; font-size: 10.5pt; line-height: 1.45; box-decoration-break: clone; }
.doc-title { font-size: 17pt; font-weight: 700; margin: 0 0 2pt; }
.doc-sub { color: #555; font-size: 9.5pt; margin-bottom: 14pt; border-bottom: 1px solid #000; padding-bottom: 6pt; }
h2 { font-size: 13pt; margin: 14pt 0 5pt; } h3 { font-size: 11.5pt; margin: 12pt 0 4pt; } h4, h5 { font-size: 10.5pt; margin: 10pt 0 3pt; }
p { margin: 0 0 7pt; } ul, ol { margin: 0 0 8pt; padding-left: 0.3in; } li { margin: 2pt 0; }
a { color: #1155cc; } code { font-family: monospace; font-size: 9.5pt; }
"""
# Word versions of the DOC_CSS headings by Markdown level ("#" prints as h2): (size, space before, space after) in points.
_DOC_HEADINGS = {1: (13, 14, 5), 2: (11.5, 12, 4), 3: (10.5, 10, 3), 4: (10.5, 10, 3)}


def markdown_doc_html(md: str, title: str = "", subtitle: str = "") -> str:
    """A generic document as one HTML page, with an optional title and a subtitle (such as the job and company)."""
    head = (f'<div class="doc-title">{html.escape(title)}</div>' if title else "") + \
           (f'<div class="doc-sub">{html.escape(subtitle)}</div>' if subtitle else "")
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title><style>{DOC_CSS}</style></head>'
            f'<body><div class="page">{head}{markdown_html(md)}</div></body></html>')


async def markdown_pdf(md: str, out: Path, title: str = "", subtitle: str = "") -> int:
    """Print a generic document to the PDF file out. Returns the page count; long documents simply take more pages."""
    out.parent.mkdir(parents=True, exist_ok=True)
    return await html_to_pdf(markdown_doc_html(md, title, subtitle), out)


def markdown_docx(md: str, out: Path, title: str = "", subtitle: str = "") -> None:
    """Write a generic document as a Word file at out, laid out like markdown_doc_html (same margins, heading sizes and
    lists). Heading text drops its ** markers and is set in bold."""
    from docx.shared import Pt
    md, title, subtitle = _xml_safe(md), _xml_safe(title), _xml_safe(subtitle)
    doc = _new_docx(top=0.75, side=0.85, bottom=0.75, font_pt=10.5)
    if title:
        run = doc.add_paragraph().add_run(title)
        run.bold, run.font.size = True, Pt(17)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        _bottom_border(p)
        p.paragraph_format.space_after = Pt(12)
    for kind, text, level in _markdown_lines(md):
        if kind == "blank":
            continue
        if kind == "heading":
            size, before, after = _DOC_HEADINGS[level]
            p = doc.add_paragraph()
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
            run = p.add_run(re.sub(r"\*\*(.+?)\*\*", r"\1", text))
            run.bold, run.font.size = True, Pt(size)
            continue
        p = doc.add_paragraph(style="List Bullet" if kind == "bullet" else None)
        _docx_runs(p, text)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
    _save(doc, out)


# ---------------------------------------------------------------- email drafts

def eml(to: str, subject: str, body: str, sender: str = "") -> bytes:
    """An outreach email as the bytes of an .eml file. The X-Unsent header makes Outlook, Thunderbird and Apple Mail open it
    as an editable draft instead of a received message. Header values are made one line (the AI writes the subject, and
    a line break in a header makes the email library refuse the whole message)."""
    def one_line(value: str) -> str:
        return " ".join((value or "").split())

    msg = EmailMessage()
    if sender:
        msg["From"] = one_line(sender)
    msg["To"] = one_line(to)
    msg["Subject"] = one_line(subject)
    msg["X-Unsent"] = "1"
    msg.set_content(body or "")
    return bytes(msg)
